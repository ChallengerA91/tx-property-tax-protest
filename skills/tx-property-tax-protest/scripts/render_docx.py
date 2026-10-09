"""protest_letter.docx: the letter sent to the CAD. It deliberately carries no disclaimer."""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import _Cell

from analysis import Analysis
from fmt import comp_when, local_address, long_date, money, number, parse_date, percent

HEADER_FILL = "D9E2F3"
DEFAULT_ADDRESSEE = "Appraisal Review Board"


def _set_font(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.05
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.8)


def _lines(document: Document, lines: Sequence[str], space_after: float = 0) -> None:
    for line in lines:
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.space_after = Pt(space_after)
    document.paragraphs[-1].paragraph_format.space_after = Pt(8)


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.keep_with_next = True


def _bullet(document: Document, text: str, bold_prefix: Optional[str] = None) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        paragraph.add_run(bold_prefix).bold = True
    paragraph.add_run(text)


def _shade(cell: _Cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _cell_text(cell: _Cell, text: str, bold: bool = False, right: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)


def _table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[str]],
           widths: Sequence[float], right_from: int = 1) -> None:
    """Compact bordered table; columns at index >= right_from are right-aligned."""
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _cell_text(cell, header, bold=True, right=index >= right_from)
        _shade(cell, HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _cell_text(cells[index], value, right=index >= right_from)
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


# --- letter sections --------------------------------------------------------

def _sender_and_addressee(document: Document, a: Analysis) -> None:
    owner, cad = a.case["owner"], a.case["cad"]
    contact = " | ".join(part for part in (owner.get("phone"), owner.get("email")) if part)
    _lines(document, [owner["name"], *owner["mailing_address"].split("\n"), *([contact] if contact else [])])
    _lines(document, [long_date(a.prepared_on) if a.prepared_on else "Date: ____________________"])
    _lines(document, [cad["name"], f"Attn: {cad.get('protest_addressee', DEFAULT_ADDRESSEE)}",
                      *cad["address"].split("\n")])


def _reference_block(document: Document, a: Analysis) -> None:
    subject = a.case["subject"]
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run(f"RE: Notice of Protest, Tax Year {a.case['tax_year']}").bold = True
    facts = [("Property", subject["address"]), ("CAD account", subject["cad_account"])]
    if a.capped:
        facts += [("CAD market value", money(a.market_value)),
                  ("Taxable (capped) appraised value", money(a.appraised_value))]
    else:
        facts.append((f"Current {a.value_term}", money(a.market_value)))
    for label, value in facts:
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(0)
        line.paragraph_format.left_indent = Inches(0.35)
        line.add_run(f"{label}: ").bold = True
        line.add_run(value)
    document.paragraphs[-1].paragraph_format.space_after = Pt(8)


def _grounds(a: Analysis) -> List[tuple]:
    basis = a.case["legal_basis"]
    grounds: List[tuple] = []
    if a.strategy in ("market_value", "both"):
        claim = "the district's market value" if a.capped else "the appraised value"
        grounds.append(("Market value", basis["market_value_cite"],
                        f"{claim} is higher than the property's market value."))
    if a.strategy in ("unequal_appraisal", "both"):
        grounds.append(("Unequal appraisal", basis["unequal_cite"],
                        "the property is appraised higher, per square foot, than comparable properties."))
    grounds += [(item["label"], item["cite"], "") for item in basis.get("additional", [])]
    return grounds


def _opening(document: Document, a: Analysis) -> None:
    cad = a.case["cad"]
    document.add_paragraph(f"Dear {cad.get('protest_addressee', DEFAULT_ADDRESSEE)}:")
    document.add_paragraph(
        f"I am the owner of the property identified above. I protest its {a.case['tax_year']} {a.value_term} "
        f"of {money(a.market_value)} and ask that it be reduced to {money(a.argued_value)}. "
        "My grounds for protest are:")
    for label, cite, reason in _grounds(a):
        _bullet(document, f" ({cite})" + (f": {reason}" if reason else ""), bold_prefix=label)


def _comp_table(document: Document, a: Analysis) -> None:
    _heading(document, "Comparable sales")
    subject = a.case["subject"]["address"]
    rows = [[local_address(c.address, subject), comp_when(c.price_type, c.sale_date),
             money(c.price), number(c.sqft), money(c.price_ppsf, cents=True), money(c.adjusted_value)]
            for c in a.sale_comps]
    rows.append(["Median", "", "", "", money(a.median_price_ppsf, cents=True), money(a.median_adjusted)])
    _table(document, ["Address", "Sale date", "Price", "Sq ft", "$ / sq ft", "Adjusted value"],
           rows, [2.5, 0.95, 0.8, 0.55, 0.75, 0.9], right_from=2)


def _neighbor_table(document: Document, a: Analysis) -> None:
    unequal = a.unequal
    subject = a.case["subject"]
    _heading(document, "Neighboring properties as appraised by the district")
    rows = [[local_address(n.address, subject["address"]), money(n.appraised_value), number(n.sqft),
             money(n.ppsf, cents=True)] for n in unequal.neighbors]
    rows.append(["Median of neighbors", "", "", money(unequal.median_ppsf, cents=True)])
    street = subject["address"].split(",")[0]
    rows.append([f"Subject property ({street})", money(a.market_value), number(subject["sqft"]),
                 money(unequal.subject_ppsf, cents=True)])
    _table(document, ["Address", "Appraised value", "Sq ft", "$ / sq ft"], rows, [3.5, 1.1, 0.7, 0.9], right_from=1)


def _income_table(document: Document, a: Analysis) -> None:
    income = a.income
    _heading(document, "Income approach")
    rows = [
        ["Potential gross rent (annual)", money(income.gross_rent)],
        [f"Vacancy and collection loss ({percent(income.vacancy_pct)})", money(-income.vacancy_loss)],
        ["Effective gross income", money(income.effective_gross_income)],
        *[[f"Expense: {e.label}", money(-e.amount)] for e in income.expenses],
        ["Net operating income", money(income.noi)],
        [f"Capitalization rate ({income.cap_rate_source})", percent(income.cap_rate_pct, 2)],
        ["Income-indicated value", money(income.indicated_value)],
    ]
    _table(document, ["Item", "Amount"], rows, [5.0, 1.4], right_from=1)


def _uses_unequal(a: Analysis) -> bool:
    return a.unequal is not None and a.strategy in ("unequal_appraisal", "both")


def _evidence_bullets(a: Analysis) -> List[str]:
    bullets: List[str] = []
    subject = a.case["subject"]
    if a.sale_comps and a.median_adjusted is not None:
        gap = a.market_value - a.median_adjusted
        bullets.append(
            f"The median adjusted value of {len(a.sale_comps)} comparable sales is {money(a.median_adjusted)}, "
            f"which is {money(abs(gap))} {'below' if gap >= 0 else 'above'} the {a.value_term}.")
    if a.other_comps:
        listed = "; ".join(f"{c.address.split(',')[0]} ({c.price_type}, {money(c.price)})" for c in a.other_comps[:3])
        bullets.append(f"Supporting context only, not sales and not used in the median: {listed}.")
    if subject.get("purchase_price"):
        when = parse_date(subject.get("purchase_date"))
        bullets.append(f"I purchased the property {'on ' + long_date(when) + ' ' if when else ''}"
                       f"for {money(subject['purchase_price'])}.")
    if _uses_unequal(a):
        u = a.unequal
        bullets.append(
            f"The property is appraised at {money(u.subject_ppsf, cents=True)} per square foot, which ranks "
            f"{u.rank} of {u.group_size} (1 = highest) among the neighbors compared; their median is "
            f"{money(u.median_ppsf, cents=True)}. At that median the property would be appraised at "
            f"{money(u.indicated_value)}.")
    if a.income_supports_reduction:
        bullets.append(f"Net operating income of {money(a.income.noi)} capitalized at "
                       f"{percent(a.income.cap_rate_pct, 2)} indicates a value of {money(a.income.indicated_value)}.")
    bullets += a.case.get("arguments") or a.case.get("market_notes", [])
    return bullets


def _closing(document: Document, a: Analysis) -> None:
    basis = a.case["legal_basis"]
    valuation = parse_date(basis.get("valuation_date"))
    as_of = f" as of {long_date(valuation)}" if valuation else ""
    _heading(document, "Opinion of value and request")
    document.add_paragraph(
        f"Based on this evidence, my opinion of the property's market value{as_of} is {money(a.argued_value)}.")
    document.add_paragraph(
        f"I request that the {a.case['tax_year']} {a.value_term} be reduced from {money(a.market_value)} to "
        f"{money(a.argued_value)}, a reduction of {money(a.reduction)} ({percent(a.reduction_pct)}). "
        "I also request the opportunity to present this evidence at a hearing.")
    document.add_paragraph("Enclosure: comparable sales analysis (comp_analysis.xlsx)")
    owner = a.case["owner"]
    closing = document.add_paragraph("Respectfully submitted,")
    closing.paragraph_format.space_before = Pt(6)
    closing.paragraph_format.keep_with_next = True
    signature = document.add_paragraph("\n______________________________")
    signature.paragraph_format.space_after = Pt(0)
    signature.paragraph_format.keep_with_next = True
    _lines(document, [owner["name"], *owner["mailing_address"].split("\n")])


def write_letter(a: Analysis, path: Path) -> None:
    document = Document()
    _set_font(document)
    _sender_and_addressee(document, a)
    _reference_block(document, a)
    _opening(document, a)
    _heading(document, "Evidence")
    for bullet in _evidence_bullets(a):
        _bullet(document, bullet)
    if a.sale_comps:
        _comp_table(document, a)
    if _uses_unequal(a):
        _neighbor_table(document, a)
    if a.income_supports_reduction:
        _income_table(document, a)
    _closing(document, a)
    _set_properties(document, a)
    document.save(path)


def _set_properties(document: Document, a: Analysis) -> None:
    props = document.core_properties
    props.author = a.case["owner"]["name"]
    props.last_modified_by = a.case["owner"]["name"]
    props.title = f"Notice of protest - {a.case['subject']['address']}"
    props.subject = f"Tax year {a.case['tax_year']}"
    props.comments = ""
    props.created = a.timestamp
    props.modified = a.timestamp
    props.revision = 1
