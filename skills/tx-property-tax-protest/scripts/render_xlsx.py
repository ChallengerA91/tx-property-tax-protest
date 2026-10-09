"""comp_analysis.xlsx: comp table, adjustments, unequal appraisal, income approach.

Every piece of user-supplied text is stored as a string cell, so text that starts with = + - or @
is shown as written and can never run as a spreadsheet formula.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Any, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from analysis import Analysis, CompRow, IncomeAnalysis, UnequalAnalysis
from fmt import MISSING, beds_baths, comp_when, long_date, money, parse_date, percent
from savings import meter_for, range_text, savings_notes, scenario_text

NAVY, BLUE, GRAY, GOLD, SKY = "1F3864", "2F5597", "F2F2F2", "FFF2CC", "DDEBF7"
MONEY = '"$"#,##0'
MONEY_CENTS = '"$"#,##0.00'
INTEGER = "#,##0"
DECIMAL = "#,##0.0"
PERCENT = '0.0"%"'
DATE_FORMAT = "yyyy-mm-dd"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


class Sheet:
    """Row-by-row writer with the formatting conventions shared by every sheet."""

    def __init__(self, ws: Worksheet, widths: Sequence[float]):
        self.ws = ws
        self.row = 1
        self.last_col = len(widths)
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        self.widths = list(widths)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "Page &P of &N"

    def _put(self, col: int, value: Any) -> Cell:
        """Write a cell. Text is forced to a string cell: openpyxl would turn '=...' into a live formula."""
        cell = self.ws.cell(self.row, col)
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"
        return cell

    def _merge(self, first_col: int, last_col: int) -> None:
        if last_col > first_col:
            self.ws.merge_cells(start_row=self.row, start_column=first_col,
                                end_row=self.row, end_column=last_col)

    def _span_chars(self, first_col: int, last_col: int) -> float:
        return sum(self.widths[first_col - 1:last_col])

    def _height_for(self, text: str, first_col: int, last_col: int, size: float = 11) -> float:
        chars_per_line = max(self._span_chars(first_col, last_col) * 11 / size * 0.95, 10)
        lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
        return max(15.0, lines * (size + 4))

    def title(self, text: str, subtitle: Optional[str] = None) -> None:
        self._merge(1, self.last_col)
        self._put(1, text).font = Font(bold=True, size=16, color=NAVY)
        self.ws.row_dimensions[self.row].height = 26
        self.row += 1
        if subtitle:
            self._merge(1, self.last_col)
            self._put(1, subtitle).font = Font(italic=True, color="595959")
            self.row += 1
        self.row += 1

    def section(self, text: str) -> None:
        self._merge(1, self.last_col)
        cell = self._put(1, text)
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(vertical="center")
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def pair(self, label: str, value: Any, fmt: Optional[str] = None, highlight: bool = False,
             value_cols: int = 3, wrap: bool = False) -> None:
        """Label in column A, value merged over the next ``value_cols`` columns."""
        label_cell = self._put(1, label)
        label_cell.font = Font(bold=True)
        label_cell.fill = _fill(GOLD if highlight else GRAY)
        label_cell.border = BORDER
        label_cell.alignment = Alignment(vertical="center", wrap_text=True)
        end = 1 + value_cols
        self._merge(2, end)
        for col in range(2, end + 1):
            self.ws.cell(self.row, col).border = BORDER
            if highlight:
                self.ws.cell(self.row, col).fill = _fill(GOLD)
        cell = self._put(2, MISSING if value is None else value)
        cell.font = Font(bold=highlight)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=wrap)
        if fmt and value is not None:
            cell.number_format = fmt
        if wrap and isinstance(value, str):
            self.ws.row_dimensions[self.row].height = self._height_for(value, 2, end)
        self.row += 1

    def header(self, labels: Sequence[str]) -> None:
        for col, label in enumerate(labels, start=1):
            cell = self._put(col, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _fill(BLUE)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[self.row].height = 32
        self.row += 1

    def data_row(self, values: Sequence[Any], formats: Sequence[Optional[str]],
                 highlight: Optional[str] = None, bold: bool = False) -> None:
        tallest = 15.0
        for col, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = self._put(col, MISSING if value is None else value)
            cell.border = BORDER
            cell.font = Font(bold=bold)
            numeric = isinstance(value, (int, float)) or hasattr(value, "isoformat")
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="right" if numeric else "left")
            if fmt and value is not None:
                cell.number_format = fmt
            if highlight:
                cell.fill = _fill(highlight)
            if isinstance(value, str):
                tallest = max(tallest, self._height_for(value, col, col))
        self.ws.row_dimensions[self.row].height = tallest
        self.row += 1

    def note(self, text: str, italic: bool = False, size: float = 11) -> None:
        self._merge(1, self.last_col)
        cell = self._put(1, text)
        cell.font = Font(italic=italic, size=size, color="404040")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.row_dimensions[self.row].height = self._height_for(text, 1, self.last_col, size)
        self.row += 1

    def gap(self) -> None:
        self.row += 1

    def footer(self, disclaimer: str) -> None:
        self.gap()
        self.note(disclaimer, italic=True, size=9)


# --- sheets -----------------------------------------------------------------

def _subject_block(sheet: Sheet, a: Analysis) -> None:
    subject = a.case["subject"]
    sheet.section("Subject property")
    sheet.pair("Address", subject["address"], value_cols=5)
    sheet.pair("CAD account", subject["cad_account"], value_cols=5)
    sheet.pair("Appraisal district", f'{a.case["cad"]["name"]} ({a.case["county"]})', value_cols=5)
    if subject.get("legal_description") or subject.get("subdivision"):
        sheet.pair("Legal description", subject.get("legal_description") or subject.get("subdivision"),
                   value_cols=5, wrap=True)
    sheet.pair("Living area (sq ft)", subject["sqft"], INTEGER, value_cols=5)
    sheet.pair("Beds / baths", beds_baths(subject.get("beds"), subject.get("baths")), value_cols=5)
    sheet.pair("Year built", subject.get("year_built"), "0", value_cols=5)
    sheet.pair("Lot size (sq ft)", subject.get("lot_sqft"), INTEGER, value_cols=5)
    if a.capped:
        sheet.pair("CAD market value", a.market_value, MONEY, value_cols=5)
        sheet.pair("Market $ / sq ft", a.market_ppsf, MONEY_CENTS, value_cols=5)
        sheet.pair("Taxable (capped) appraised value", a.appraised_value, MONEY, value_cols=5)
    else:
        sheet.pair("Appraised value", a.market_value, MONEY, value_cols=5)
        sheet.pair("Appraised $ / sq ft", a.market_ppsf, MONEY_CENTS, value_cols=5)
    if subject.get("purchase_price"):
        bought = long_date(parse_date(subject.get("purchase_date"))) if subject.get("purchase_date") else "date not given"
        sheet.pair("Purchase price", subject["purchase_price"], MONEY, value_cols=5)
        sheet.pair("Purchase date", bought, value_cols=5)


COMP_HEADERS = ["Address", "Sale date", "Price type", "Price", "Sq ft", "$ / sq ft", "Beds / baths",
                "Year built", "Distance (mi)", "Net adjustments", "Adjusted value", "Adj. $ / sq ft",
                "Source", "Notes"]
COMP_FORMATS = [None, DATE_FORMAT, None, MONEY, INTEGER, MONEY_CENTS, None, "0", DECIMAL, MONEY, MONEY,
                MONEY_CENTS, None, None]


def _comp_values(c: CompRow) -> List[Any]:
    is_sale = c.price_type == "sale"
    kind = "Sale" if is_sale else comp_when(c.price_type, c.sale_date)    # a listing date is not a sale date
    return [c.address, c.sale_date if is_sale else None, kind, c.price, c.sqft, c.price_ppsf,
            beds_baths(c.beds, c.baths), c.year_built, c.distance_mi, c.adjustment_total,
            c.adjusted_value, c.adjusted_ppsf, c.source, c.notes]


def _comp_sheet(wb: Workbook, a: Analysis) -> None:
    ws = wb.active
    ws.title = "Comp Analysis"
    sheet = Sheet(ws, [36, 12, 11, 14, 9, 11, 11, 9, 10, 14, 15, 12, 16, 42])
    case = a.case
    prepared = f" | Prepared {long_date(a.prepared_on)}" if a.prepared_on else ""
    sheet.title(f"Comparable Sales Analysis - Tax Year {case['tax_year']}",
                f"{case['cad']['name']}{prepared}")
    _subject_block(sheet, a)
    sheet.gap()
    sheet.section("Comparable sales")
    sheet.header(COMP_HEADERS)
    if not a.sale_comps:
        sheet.note("No comparable sales were supplied for this case.")
    for comp in a.sale_comps:
        sheet.data_row(_comp_values(comp), COMP_FORMATS)
    if a.other_comps:
        sheet.gap()
        sheet.section("Other data points (not used in the median)")
        sheet.header(COMP_HEADERS)
        for comp in a.other_comps:
            sheet.data_row(_comp_values(comp), COMP_FORMATS)
    sheet.gap()
    _summary_block(sheet, a)
    sheet.gap()
    _notes_block(sheet, a)
    sheet.footer(a.disclaimer)


def _summary_block(sheet: Sheet, a: Analysis) -> None:
    sheet.section("Summary")
    sheet.pair("Comparable sales used", len(a.sale_comps), "0")
    sheet.pair("Median adjusted value", a.median_adjusted, MONEY, highlight=True)
    sheet.pair("Average adjusted value", a.mean_adjusted, MONEY, highlight=True)
    sheet.pair("Median $ / sq ft (unadjusted)", a.median_price_ppsf, MONEY_CENTS)
    sheet.pair("Median adjusted $ / sq ft", a.median_adjusted_ppsf, MONEY_CENTS)
    sheet.pair(f"CAD {a.value_term}", a.market_value, MONEY)
    if a.capped:
        sheet.pair("Taxable (capped) appraised value", a.appraised_value, MONEY)
    sheet.pair("Argued market value", a.argued_value, MONEY, highlight=True)
    sheet.pair("Argued $ / sq ft", a.argued_ppsf, MONEY_CENTS)
    sheet.pair("Reduction in market value" if a.capped else "Potential reduction", a.reduction, MONEY, highlight=True)
    sheet.pair("Reduction as % of market value" if a.capped else "Reduction as % of appraisal",
               a.reduction_pct, PERCENT)
    if a.capped:
        sheet.pair("Reduction below the taxable (capped) value", a.savings.reduction, MONEY)
    sheet.pair("Annual tax savings from protest",
               "$0" if a.savings.protest_blocked else range_text(a.savings.protest.low, a.savings.protest.high))
    sheet.pair("Savings meter", meter_for(a.savings), value_cols=5, wrap=bool(savings_notes(a.savings)))


def _notes_block(sheet: Sheet, a: Analysis) -> None:
    case = a.case
    sheet.section("Notes")
    sheet.note("Adjusted value = price + net adjustments (see the Adjustments sheet). A positive "
               "adjustment means the comparable is inferior to the subject; a negative one means it is superior.")
    if a.other_comps:
        sheet.note("Listings and estimates are shown separately as supporting data points; "
                   "only sales enter the median and average.")
    rates = case["tax_rates"]
    source = f" Source: {rates['source']}." if rates.get("source") else ""
    sheet.note(f"Savings use a total tax rate of {rates['total']:g} per $100 of value.{source} "
               "Low and high are settlement scenarios, not predictions "
               f"({scenario_text(a.savings)}).")
    for note in savings_notes(a.savings):
        sheet.note(note)
    for note in case.get("market_notes", []):
        sheet.note(f"- {note}")


def _adjustments_sheet(wb: Workbook, a: Analysis) -> None:
    sheet = Sheet(wb.create_sheet("Adjustments"), [40, 46, 16])
    sheet.title("Comparable Adjustments", "Dollar adjustments applied to each comparable price")
    sheet.header(["Comparable", "Adjustment", "Amount"])
    for comp in [*a.sale_comps, *a.other_comps]:
        price_label = "Price" if comp.price_type == "sale" else f"Price ({comp.price_type}, not used in the median)"
        sheet.data_row([comp.address, price_label, comp.price], [None, None, MONEY], highlight=SKY, bold=True)
        for adj in comp.adjustments:
            sheet.data_row([None, adj.label, adj.amount], [None, None, MONEY])
        if not comp.adjustments:
            sheet.data_row([None, "No adjustments", 0], [None, None, MONEY])
        sheet.data_row([None, "Adjusted value", comp.adjusted_value], [None, None, MONEY],
                       highlight=GOLD, bold=True)
    sheet.footer(a.disclaimer)


def _unequal_sheet(wb: Workbook, a: Analysis, u: UnequalAnalysis) -> None:
    sheet = Sheet(wb.create_sheet("Unequal Appraisal"), [8, 40, 16, 10, 12, 11, 16])
    basis = "market value" if a.capped else "appraised value"
    sheet.title("Unequal Appraisal Analysis",
                f"Neighboring properties as appraised by the CAD, compared on {basis} $ per square foot")
    sheet.header(["Rank", "Address", "Appraised value", "Sq ft", "$ / sq ft", "Year built", "Subject minus this"])
    for row in a.ranked_neighbors():
        gap = None if row.is_subject else u.subject_ppsf - row.ppsf
        sheet.data_row([row.rank, row.address, row.appraised_value, row.sqft, row.ppsf, row.year_built, gap],
                       ["0", None, MONEY, INTEGER, MONEY_CENTS, "0", MONEY_CENTS],
                       highlight=GOLD if row.is_subject else None, bold=row.is_subject)
    sheet.gap()
    sheet.section("Summary")
    sheet.pair("Neighbors compared", len(u.neighbors), "0", value_cols=3)
    sheet.pair("Neighbor median $ / sq ft", u.median_ppsf, MONEY_CENTS, value_cols=3)
    sheet.pair("Neighbor average $ / sq ft", u.mean_ppsf, MONEY_CENTS, value_cols=3)
    sheet.pair("Neighbor range $ / sq ft", f"{money(u.low_ppsf, True)} to {money(u.high_ppsf, True)}", value_cols=3)
    sheet.pair("Subject market value $ / sq ft" if a.capped else "Subject $ / sq ft", u.subject_ppsf, MONEY_CENTS,
               highlight=True, value_cols=3)
    sheet.pair("Subject rank (1 = highest)", f"{u.rank} of {u.group_size}", highlight=True, value_cols=3)
    sheet.pair("Subject percentile", percent(u.percentile, 0), value_cols=3)
    sheet.pair("Equalized value (median x sq ft)", u.indicated_value, MONEY, highlight=True, value_cols=3)
    sheet.pair("CAD market value" if a.capped else "Appraised value", a.market_value, MONEY, value_cols=3)
    sheet.pair("Difference", a.market_value - u.indicated_value, MONEY, value_cols=3)
    sheet.note("Percentile = share of neighbors appraised at or below the subject's $ per sq ft; ties count half. "
               "Equal $ per sq ft values share a rank.")
    if a.capped:
        sheet.note("The subject is compared at the CAD market value, not the capped appraised value.")
    sheet.footer(a.disclaimer)


def _income_sheet(wb: Workbook, a: Analysis, inc: IncomeAnalysis) -> None:
    sheet = Sheet(wb.create_sheet("Income Approach"), [44, 18, 16, 40])
    sheet.title("Income Approach", "Net operating income divided by the capitalization rate")
    sheet.header(["Line", "Amount", "Rate", "Source / note"])
    sheet.data_row(["Potential gross rent (annual)", inc.gross_rent, None, inc.rent_source],
                   [None, MONEY, None, None])
    sheet.data_row(["Less: vacancy and collection loss", -inc.vacancy_loss, inc.vacancy_pct, None],
                   [None, MONEY, PERCENT, None])
    sheet.data_row(["Effective gross income", inc.effective_gross_income, None, None],
                   [None, MONEY, None, None], highlight=GRAY, bold=True)
    for expense in inc.expenses:
        sheet.data_row([f"Less: {expense.label}", -expense.amount, None, None], [None, MONEY, None, None])
    sheet.data_row(["Total operating expenses", -inc.total_expenses, None, None],
                   [None, MONEY, None, None], highlight=GRAY, bold=True)
    sheet.data_row(["Net operating income (NOI)", inc.noi, None, None],
                   [None, MONEY, None, None], highlight=GOLD, bold=True)
    sheet.data_row(["Capitalization rate", None, inc.cap_rate_pct, inc.cap_rate_source],
                   [None, None, PERCENT, None])
    sheet.data_row(["Income-indicated value (NOI / cap rate)", inc.indicated_value, None, None],
                   [None, MONEY, None, None], highlight=GOLD, bold=True)
    term = "market" if a.capped else "appraised"
    sheet.data_row([f"CAD {a.value_term}", a.market_value, None, None], [None, MONEY, None, None])
    sheet.data_row([f"Difference ({term} minus indicated)", a.market_value - inc.indicated_value,
                    None, None], [None, MONEY, None, None])
    if not a.income_supports_reduction:
        sheet.gap()
        sheet.note("The income approach indicates a value at or above the CAD value, so the protest letter "
                   "does not use it.")
    if inc.circuit_breaker_note:
        sheet.gap()
        sheet.section("Appraisal-increase limit (circuit breaker)")
        sheet.note(inc.circuit_breaker_note)
    sheet.footer(a.disclaimer)


# --- entry point ------------------------------------------------------------

def write_xlsx(a: Analysis, path: Path) -> None:
    wb = Workbook()
    _comp_sheet(wb, a)
    _adjustments_sheet(wb, a)
    if a.unequal:
        _unequal_sheet(wb, a, a.unequal)
    if a.income:
        _income_sheet(wb, a, a.income)
    wb.properties.creator = "tx-property-tax-protest"
    wb.properties.lastModifiedBy = "tx-property-tax-protest"
    wb.properties.title = f"Comp analysis - {a.case['subject']['address']}"
    wb.properties.created = a.timestamp
    wb.properties.modified = a.timestamp
    wb.save(path)
