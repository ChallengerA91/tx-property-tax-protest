"""Second and third review batches: school ceiling, optional rental income, listings, warnings, deadline inputs, holidays."""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys

import pytest
from openpyxl import load_workbook

import build_package
import county_lookup
import deadlines
import savings
from analysis import build_analysis
from case_schema import case_warnings, load_case, validate_case
from conftest import SCRIPTS
from fmt import comp_when, strict_date
from test_build_package import find_row, letter_text, sheet_text

D = strict_date


def has_error(errors, fragment):
    return any(fragment in error for error in errors)


def run_cli(*args):
    return subprocess.run([sys.executable, str(SCRIPTS / "build_package.py"), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8")


def build(case, tmp_path, write_case, name="out"):
    out = tmp_path / name
    build_package.build_package(write_case(case, f"{name}.json"), out)
    return out


# --- school tax ceiling -----------------------------------------------------------------------------

CEILING_CASE = {
    "subject": {"appraised_value": 400000}, "argued_value": 350000,
    "tax_rates": {"total": 2.0, "school": 1.2, "school_ceiling": True},
    "exemptions": {"eligible": [{"name": "School", "amount": 100000, "applies_to": "school"},
                                {"name": "Local", "amount": 10000, "applies_to": "all"}]},
}


def test_school_ceiling_counts_no_school_tax_savings():
    estimate = savings.estimate_savings(CEILING_CASE)
    assert estimate.school_ceiling
    assert (estimate.protest.low, estimate.protest.high) == (pytest.approx(200.0), pytest.approx(400.0))  # 0.8 per $100
    assert {e.name: e.annual for e in estimate.exemptions} == {"School": 0.0, "Local": pytest.approx(80.0)}
    assert (estimate.total.low, estimate.total.high) == (pytest.approx(280.0), pytest.approx(480.0))
    assert savings.meter_for(estimate).endswith(savings.SCHOOL_CEILING_NOTE)
    assert savings.SCHOOL_CEILING_NOTE == ("A school tax ceiling applies, so school-tax savings are not counted; "
                                           "the real figure depends on the ceiling amount.")


def test_without_the_ceiling_the_same_case_counts_school_tax():
    case = copy.deepcopy(CEILING_CASE)
    case["tax_rates"]["school_ceiling"] = False
    estimate = savings.estimate_savings(case)
    assert estimate.protest.high == pytest.approx(1000.0) and not estimate.school_ceiling
    assert savings.SCHOOL_CEILING_NOTE not in savings.meter_for(estimate)


def test_school_ceiling_needs_the_school_rate(homestead_case):
    del homestead_case["tax_rates"]["school"]
    homestead_case["exemptions"]["eligible"] = []
    homestead_case["tax_rates"]["school_ceiling"] = True
    assert has_error(validate_case(homestead_case), "tax_rates.school: required when tax_rates.school_ceiling is true")
    with pytest.raises(ValueError, match="school_ceiling"):
        savings.estimate_savings({"subject": {"appraised_value": 10}, "tax_rates": {"total": 2, "school_ceiling": True}})


def test_school_ceiling_must_be_a_boolean(homestead_case):
    homestead_case["tax_rates"]["school_ceiling"] = "yes"
    assert has_error(validate_case(homestead_case), "tax_rates.school_ceiling: must be true or false")


def test_every_savings_output_carries_the_ceiling_note(tmp_path, write_case, homestead_case):
    homestead_case["tax_rates"]["school_ceiling"] = True
    out = build(homestead_case, tmp_path, write_case)
    note = savings.SCHOOL_CEILING_NOTE
    sheet = load_workbook(out / "comp_analysis.xlsx")["Comp Analysis"]
    assert find_row(sheet, "Savings meter")[1].value == f"Estimated savings so far: $220 to $440 per year. {note}"
    assert note in sheet_text(sheet)
    checklist = (out / "filing_checklist.md").read_text(encoding="utf-8")
    assert "**Estimated savings so far: $220 to $440 per year**" in checklist and checklist.count(note) == 2
    assert "| Exemption: Over-65 or disabled additional school exemption | $0 | $0 |" in checklist
    html = (out / "dashboard.html").read_text(encoding="utf-8")
    assert html.count(note) == 2
    result = run_cli(write_case(homestead_case, "cli.json"), "--out", tmp_path / "cli")
    assert note in result.stdout


# --- rental income is optional --------------------------------------------------------------------

def test_rental_without_income_leaves_out_every_income_output(tmp_path, write_case, rental_case):
    del rental_case["income"]
    out = build(rental_case, tmp_path, write_case)
    assert load_workbook(out / "comp_analysis.xlsx").sheetnames == ["Comp Analysis", "Adjustments", "Unequal Appraisal"]
    assert "Income" not in letter_text(out)
    assert "Income" not in (out / "dashboard.html").read_text(encoding="utf-8")
    assert "Income-indicated" not in (out / "filing_checklist.md").read_text(encoding="utf-8")


def test_a_rental_needs_some_evidence(rental_case):
    del rental_case["income"], rental_case["neighbors"]
    rental_case["comps"] = []
    assert has_error(validate_case(rental_case), "rental cases need at least one of")


def test_a_rental_can_lead_with_neighbors_alone(rental_case):
    del rental_case["income"]
    rental_case["comps"] = []
    rental_case["strategy"] = "unequal_appraisal"
    rental_case["legal_basis"]["unequal_cite"] = "Tax Code §41.41(a)(2)"
    assert validate_case(rental_case) == []


def test_income_present_still_needs_a_cap_rate_source(rental_case):
    rental_case["income"]["cap_rate_source"] = "  "
    assert has_error(validate_case(rental_case), "income.cap_rate_source: must be a non-empty string")


def test_income_approach_is_left_out_of_the_letter_when_it_does_not_lower_the_value(tmp_path, write_case, rental_case):
    rental_case["income"]["gross_rent"] = 60000          # indicated value now far above the CAD value
    out = build(rental_case, tmp_path, write_case)
    a = build_analysis(load_case(write_case(rental_case, "check.json")))
    assert a.income.indicated_value > a.market_value and not a.income_supports_reduction
    text = letter_text(out)
    assert "Income approach" not in text and "Net operating income" not in text
    assert "Income Approach" in load_workbook(out / "comp_analysis.xlsx").sheetnames
    assert "Income approach" in (out / "dashboard.html").read_text(encoding="utf-8")
    assert "does not use it" in sheet_text(load_workbook(out / "comp_analysis.xlsx")["Income Approach"])


def test_the_cli_warns_when_the_income_approach_is_not_used(tmp_path, write_case, rental_case):
    rental_case["income"]["gross_rent"] = 60000
    result = run_cli(write_case(rental_case), "--out", tmp_path / "out")
    assert result.returncode == 0
    assert re.search(r"warning: Income approach indicates \$[\d,]+, above the CAD value; it is not used in the letter",
                     result.stderr)


def test_the_income_approach_stays_in_the_letter_when_it_is_below_the_cad_value(tmp_path, write_case, rental_case):
    out = build(rental_case, tmp_path, write_case)
    assert "Income approach" in letter_text(out)
    assert "warning" not in run_cli(write_case(rental_case, "c.json"), "--out", tmp_path / "o2").stderr


# --- listings and estimates are not sales ------------------------------------------------------------

def test_comp_when_never_shows_a_bare_date_for_a_listing():
    assert comp_when("sale", D("2026-09-18")) == "Sep 18, 2026"
    assert comp_when("sale", None) == "Sale (date not given)"
    assert comp_when("listing", None) == "Listing"
    assert comp_when("listing", D("2026-12-01")) == "Listing, listed Dec 1, 2026"
    assert comp_when("estimate", D("2027-01-01")) == "Estimate, as of Jan 1, 2027"


def test_listings_stay_out_of_the_median_and_get_their_own_table(tmp_path, write_case, homestead_case):
    homestead_case["comps"][-1]["sale_date"] = "2026-12-01"           # a listing date, not a sale date
    out = build(homestead_case, tmp_path, write_case)
    book = load_workbook(out / "comp_analysis.xlsx")
    sheet = book["Comp Analysis"]
    assert find_row(sheet, "Comparable sales used")[1].value == 5
    assert find_row(sheet, "Median adjusted value")[1].value == 421000
    assert any(row[0].value == "Other data points (not used in the median)" for row in sheet.iter_rows())
    listing = next(row for row in sheet.iter_rows() if row[0].value and str(row[0].value).startswith("14 Mock Orchard"))
    assert listing[1].value == "—" and listing[2].value == "Listing, listed Dec 1, 2026"
    html = (out / "dashboard.html").read_text(encoding="utf-8")
    assert "Other data points (not used in the median)" in html and "Listing, listed Dec 1, 2026" in html
    assert "<td>Dec 1, 2026</td>" not in html
    checklist = load_workbook(out / "comp_analysis.xlsx")["Adjustments"]
    assert "not used in the median" in sheet_text(checklist)


def test_the_letter_counts_only_sales_and_mentions_listings_as_context(tmp_path, write_case, homestead_case):
    out = build(homestead_case, tmp_path, write_case)
    text = letter_text(out)
    assert "The median adjusted value of 5 comparable sales is $421,000" in text
    assert "Supporting context only, not sales and not used in the median: 14 Mock Orchard Road (listing, $455,000)." in text
    from docx import Document
    tables = "\n".join(cell.text for table in Document(out / "protest_letter.docx").tables
                       for row in table.rows for cell in row.cells)
    assert "14 Mock Orchard" not in tables and "Sep 18, 2026" in tables


def test_a_market_value_strategy_needs_at_least_one_sale(homestead_case):
    for comp in homestead_case["comps"]:
        comp["price_type"] = "listing"
    errors = validate_case(homestead_case)
    assert has_error(errors, "comps: at least one comparable with price_type 'sale' is required")
    assert has_error(errors, "listings and estimates are not used in the median")


# --- warnings and cross-checks --------------------------------------------------------------------------

def test_the_samples_raise_no_warnings(homestead_case, rental_case, capped_case):
    for case in (homestead_case, rental_case, capped_case):
        assert case_warnings(case) == []


def test_hearing_and_order_dates_before_the_notice_warn(homestead_case):
    homestead_case["hearing_date"] = "2027-03-01"
    homestead_case["arb_order_date"] = "2027-02-01"
    warnings = case_warnings(homestead_case)
    assert "hearing_date 2027-03-01 is earlier than notice_date 2027-04-01" in warnings
    assert "arb_order_date 2027-02-01 is earlier than notice_date 2027-04-01" in warnings


def test_a_sale_after_prepared_on_warns(homestead_case):
    homestead_case["comps"][0]["sale_date"] = "2027-05-01"
    assert case_warnings(homestead_case) == ["comps[0].sale_date 2027-05-01 is later than prepared_on (2027-04-20)"]


def test_without_prepared_on_the_latest_case_date_is_the_yardstick_not_the_clock(homestead_case):
    del homestead_case["prepared_on"]
    homestead_case["hearing_date"] = "2027-06-15"
    homestead_case["comps"][0]["sale_date"] = "2027-07-01"
    assert case_warnings(homestead_case) == [
        "comps[0].sale_date 2027-07-01 is later than the latest date in the case (2027-06-15)"]


def test_a_sale_more_than_five_years_before_the_valuation_date_warns(homestead_case):
    homestead_case["comps"][0]["sale_date"] = "2021-12-31"
    homestead_case["comps"][1]["sale_date"] = "2022-01-01"            # exactly five years: fine
    warnings = case_warnings(homestead_case)
    assert warnings == ["comps[0].sale_date 2021-12-31 is more than 5 years before the valuation date 2027-01-01"]


def test_the_build_prints_warnings_but_still_builds(tmp_path, write_case, homestead_case):
    homestead_case["comps"][0]["sale_date"] = "2027-05-01"
    result = run_cli(write_case(homestead_case), "--out", tmp_path / "out")
    assert result.returncode == 0
    assert "warning: comps[0].sale_date 2027-05-01 is later than prepared_on (2027-04-20)" in result.stderr
    assert (tmp_path / "out" / "comp_analysis.xlsx").exists()


# --- capped homestead labels ------------------------------------------------------------------------------

def test_capped_unequal_appraisal_is_labeled_and_ranked_on_market_value(tmp_path, write_case, capped_case):
    capped_case["strategy"] = "both"
    capped_case["legal_basis"]["unequal_cite"] = "Tax Code §41.41(a)(2)"
    capped_case["neighbors"] = [{"address": "N1", "appraised_value": 500000, "sqft": 2400},
                                {"address": "N2", "appraised_value": 480000, "sqft": 2400}]
    out = build(capped_case, tmp_path, write_case)
    a = build_analysis(load_case(write_case(capped_case, "a.json")))
    assert a.unequal.subject_ppsf == pytest.approx(520000 / 2400) and a.unequal.rank == 1   # 462,000 would rank last
    sheet = load_workbook(out / "comp_analysis.xlsx")["Unequal Appraisal"]
    assert find_row(sheet, "Subject market value $ / sq ft")[1].value == pytest.approx(520000 / 2400)
    assert find_row(sheet, "Subject rank (1 = highest)")[1].value == "1 of 3"
    assert "not the capped appraised value" in sheet_text(sheet)
    assert "Subject market-value $/sq ft rank among neighbors (1 = highest) | 1 of 3" in \
        (out / "filing_checklist.md").read_text(encoding="utf-8")
    html = (out / "dashboard.html").read_text(encoding="utf-8")
    assert "Market-value $ per sq ft: neighbors vs. your property" in html


def test_tied_neighbors_share_a_rank_and_the_stated_rank_matches_the_table(tmp_path, write_case, homestead_case):
    same = {"appraised_value": 465000, "sqft": 2100}                  # exactly the subject's $/sq ft
    homestead_case["neighbors"] = [{"address": "T1", **same}, {"address": "T2", **same},
                                   {"address": "L1", "appraised_value": 400000, "sqft": 2100}]
    out = build(homestead_case, tmp_path, write_case)
    a = build_analysis(load_case(write_case(homestead_case, "ties.json")))
    rows = a.ranked_neighbors()
    assert [r.rank for r in rows] == [1, 1, 1, 4]
    assert next(r for r in rows if r.is_subject).rank == a.unequal.rank == 1
    sheet = load_workbook(out / "comp_analysis.xlsx")["Unequal Appraisal"]
    table = [(row[0].value, row[1].value) for row in sheet.iter_rows(min_row=5, max_row=8)]
    assert [rank for rank, _ in table] == [1, 1, 1, 4]
    assert find_row(sheet, "Subject rank (1 = highest)")[1].value == "1 of 4"
    html = (out / "dashboard.html").read_text(encoding="utf-8")
    assert html.count("<th scope=\"row\">1. ") == 3 and "<th scope=\"row\">4. L1" in html


# --- deadline wording and the portal warning -----------------------------------------------------------------

def test_the_records_approval_date_is_not_presented_as_an_owner_deadline():
    found = {d.id: d for d in deadlines.compute_deadlines(deadlines.load_rules(), D("2027-04-01"), 2027)}
    approval = found["protest_late_good_cause"]
    assert approval.event == "ARB records-approval date (not an owner deadline)"
    assert approval.statute == "Tax Code §41.44(b); §41.12(a), (c)"
    assert "only until the ARB actually approves the records" in approval.note
    events = [r["event"] for r in json.loads((SCRIPTS / "deadline_rules.json").read_text(encoding="utf-8"))["rules"]]
    assert not [e for e in events if "must approve" in e]


def test_a_deadline_that_moved_forward_also_shows_the_last_business_day_before_it(tmp_path, write_case, homestead_case):
    out = build(homestead_case, tmp_path, write_case)
    checklist = (out / "filing_checklist.md").read_text(encoding="utf-8")
    assert ("Some county portals close on the original date; file by the last business day before it to be safe:"
            in checklist)
    assert "- Protest filing deadline: Fri, May 14, 2027 (statutory date Sat, May 15, 2027, moves to Mon, May 17, 2027)" \
        in checklist
    ics = "".join(deadlines.dt.__name__ and (out / "deadlines.ics").read_bytes().decode("utf-8").split("\r\n "))
    assert "Some county portals close on the original date" in ics


def test_a_date_that_did_not_move_has_no_portal_warning(tmp_path, write_case, rental_case):
    out = build(rental_case, tmp_path, write_case)                      # protest deadline Monday May 24, 2027
    assert "Some county portals close" not in (out / "filing_checklist.md").read_text(encoding="utf-8")


# --- deadlines.py: optional inputs ---------------------------------------------------------------------------

def test_the_tax_year_defaults_to_the_year_of_the_date_given(capsys):
    assert deadlines.main(["--arb-order-date", "2027-07-20"]) == 0
    out = capsys.readouterr().out
    assert "Deadlines for tax year 2027 (notice date not given)" in out
    assert "Fri 2027-09-17  File a court petition for review by" in out
    assert deadlines.main(["--tax-year", "2027", "--arb-order-date", "2027-07-20"]) == 0
    assert "tax year 2027" in capsys.readouterr().out
    assert deadlines.main(["--hearing-date", "2027-06-15"]) == 0
    assert "Fri 2027-05-28  ARB mails the hearing notice by" in capsys.readouterr().out


def test_skipped_rules_are_named_with_the_flag_that_would_add_them(capsys):
    deadlines.main(["--arb-order-date", "2027-07-20"])
    out = capsys.readouterr().out
    assert "Skipped, needs --notice-date: Protest filing deadline" in out
    assert "Skipped, needs --hearing-date: ARB mails the hearing notice by" in out
    deadlines.main(["--notice-date", "2027-04-01", "--hearing-date", "2027-06-15"])
    again = capsys.readouterr().out
    assert "needs --hearing-date" not in again and "Skipped, needs --arb-order-date: Request binding arbitration by" in again


def test_the_cli_needs_some_date_or_a_tax_year(capsys):
    with pytest.raises(SystemExit) as caught:
        deadlines.main([])
    assert caught.value.code == 2
    assert "give --tax-year, or at least one of" in capsys.readouterr().err


def test_a_tax_year_that_contradicts_the_notice_date_is_an_error(capsys):
    with pytest.raises(SystemExit):
        deadlines.main(["--notice-date", "2026-04-15", "--tax-year", "2027"])
    assert "does not match the year of --notice-date 2026-04-15" in capsys.readouterr().err


@pytest.mark.parametrize("text", ["20270415", "2027-4-15"])
def test_the_cli_takes_dates_only_as_yyyy_mm_dd(text, capsys):
    with pytest.raises(SystemExit):
        deadlines.main(["--notice-date", text])
    assert "YYYY-MM-DD" in capsys.readouterr().err


def test_strict_date_accepts_only_the_plain_form():
    assert strict_date("2027-04-15").isoformat() == "2027-04-15"
    for bad in ("20270415", "2027-W15-1", "2027-04-15T01:00", " 2027-04-15", None, 20270415):
        with pytest.raises(ValueError):
            strict_date(bad)


# --- deadlines: exemption dates and rule notes ------------------------------------------------------------------

def seed(notice, **inputs):
    given = {name: D(value) for name, value in inputs.items()}
    return {d.id: d for d in deadlines.compute_deadlines(deadlines.load_rules(), D(notice) if notice else None, 2027, given)}


def test_new_buyer_deadline_is_the_day_before_the_first_anniversary():
    found = seed("2027-04-01", acquisition_date="2027-03-15")
    assert found["exemption_deadline_new_buyer"].date == D("2028-03-14")
    assert found["exemption_deadline_new_buyer"].statute == "Tax Code §11.42(f); §11.43(d)"
    assert seed("2027-04-01", acquisition_date="2028-02-29")["exemption_deadline_new_buyer"].date == D("2029-02-27")


def test_age_65_or_disabled_deadline_is_the_first_anniversary_itself():
    found = seed("2027-04-01", qualified_date="2027-06-01")
    assert found["exemption_deadline_age65_or_disabled"].date == D("2028-06-01")
    assert found["exemption_deadline_age65_or_disabled"].statute == "Tax Code §11.43(k), (m)"


def test_these_rules_wait_for_their_dates():
    assert "exemption_deadline_new_buyer" not in seed("2027-04-01")
    assert "exemption_deadline_age65_or_disabled" not in seed("2027-04-01")


def test_cli_accepts_the_acquisition_and_qualified_dates(capsys):
    assert deadlines.main(["--notice-date", "2027-04-01", "--acquisition-date", "2027-03-15",
                           "--qualified-date", "2027-06-01", "--all"]) == 0
    out = capsys.readouterr().out
    assert "Tue 2028-03-14  Last day to apply for the part-year homestead exemption (new buyer)" in out
    assert "Thu 2028-06-01  Last day to apply for the 65+ or disabled exemption" in out


def test_late_homestead_note_explains_the_form_wording():
    note = seed("2027-04-01")["exemption_late_homestead"].note
    assert ("Form 50-114 says two years after the filing deadline (April 30, 2029 for tax year 2027); the statute "
            "says two years after the delinquency date. File as early as you can.") in note


def test_before_delinquency_rules_use_the_day_before():
    found = seed("2027-04-01")
    assert found["delinquency_date"].date == D("2028-02-01")
    for rule_id in ("protest_late_no_notice", "protest_late_offshore_or_deployed", "motion_to_correct_overappraisal",
                    "court_prepayment", "arbitration_prepayment", "soah_prepayment"):
        assert found[rule_id].date == D("2028-01-31"), rule_id


# --- holidays: Gov't Code 662.003 (a) and (b), computed per year ----------------------------------------------

HOLIDAYS = [
    # name, 2027, 2028, citation
    ("New Year's Day", "2027-01-01", "2028-01-01", "(a)"),
    ("Martin Luther King, Jr., Day", "2027-01-18", "2028-01-17", "(a)"),
    ("Presidents' Day", "2027-02-15", "2028-02-21", "(a)"),
    ("Memorial Day", "2027-05-31", "2028-05-29", "(a)"),
    ("Independence Day", "2027-07-04", "2028-07-04", "(a)"),
    ("Labor Day", "2027-09-06", "2028-09-04", "(a)"),
    ("Veterans Day", "2027-11-11", "2028-11-11", "(a)"),
    ("Thanksgiving Day", "2027-11-25", "2028-11-23", "(a)"),
    ("Christmas Day", "2027-12-25", "2028-12-25", "(a)"),
    ("Confederate Heroes Day", "2027-01-19", "2028-01-19", "(b)"),
    ("Texas Independence Day", "2027-03-02", "2028-03-02", "(b)"),
    ("San Jacinto Day", "2027-04-21", "2028-04-21", "(b)"),
    ("Emancipation Day in Texas", "2027-06-19", "2028-06-19", "(b)"),
    ("Lyndon Baines Johnson Day", "2027-08-27", "2028-08-27", "(b)"),
    ("Friday after Thanksgiving Day", "2027-11-26", "2028-11-24", "(b)"),
    ("December 24 (state holiday)", "2027-12-24", "2028-12-24", "(b)"),
    ("December 26 (state holiday)", "2027-12-26", "2028-12-26", "(b)"),
]


def holiday_named(name):
    return next(h for h in deadlines.load_rules().holidays if h.name == name)


def test_the_file_models_all_seventeen_days_with_their_citations():
    data = json.loads((SCRIPTS / "deadline_rules.json").read_text(encoding="utf-8"))
    assert len(data["holidays"]) == 17 == len(HOLIDAYS)
    for entry, (name, _, _, part) in zip(data["holidays"], HOLIDAYS):
        assert entry["name"] == name and entry["statute"] == f"Gov't Code §662.003{part}"
    assert "does not say whether a holiday that falls on a weekend shifts" in data["holidays_note"]
    assert "legal state or national holiday" in data["holidays_note"]


@pytest.mark.parametrize("name, day_2027, day_2028, part", HOLIDAYS)
def test_each_holiday_lands_on_the_right_day_in_2027_and_2028(name, day_2027, day_2028, part):
    holiday = holiday_named(name)
    assert holiday.date_in(2027) == D(day_2027)
    assert holiday.date_in(2028) == D(day_2028)
    assert holiday.falls_on(D(day_2027)) and not holiday.falls_on(D(day_2027).replace(day=D(day_2027).day - 1 or 2))


def test_a_weekend_holiday_is_not_shifted():
    ruleset = deadlines.load_rules()
    assert ruleset.closed_reason(D("2027-07-04")) == "Sunday"            # Independence Day on a Sunday
    assert ruleset.closed_reason(D("2027-07-05")) is None                # the Monday after stays open
    assert ruleset.closed_reason(D("2027-07-02")) is None
    assert ruleset.closed_reason(D("2027-06-18")) is None                # Juneteenth fell on Saturday the 19th


def test_labor_day_rolls_a_court_deadline():
    found = seed("2027-04-01", arb_order_date="2027-07-08")["court_petition"]   # 60 days after July 8 is Labor Day
    assert found.computed_date == D("2027-09-06") and found.date == D("2027-09-03")
    assert found.roll_reasons == ("Labor Day",)
    assert "Labor Day" in found.roll_text


def test_a_state_holiday_rolls_a_filing_forward():
    ruleset = deadlines.parse_rules({"rules": [
        {"id": "x", "event": "X", "type": "fixed_date", "params": {"month": 4, "day": 21}, "weekend_roll": True}],
        "holidays": [{"name": "San Jacinto Day", "type": "annual", "month": 4, "day": 21}]})
    result = deadlines.compute_deadlines(ruleset, None, 2027)[0]          # April 21, 2027 is a Wednesday
    assert result.date == D("2027-04-22") and result.roll_reasons == ("San Jacinto Day",)


def test_holiday_definitions_in_a_rules_file_are_validated():
    base = {"rules": []}
    for bad in ({"name": "X", "type": "annual", "month": 2, "day": 30},
                {"name": "X", "type": "after", "holiday": "Missing", "days": 1},
                {"name": "X", "type": "after", "holiday": "X", "days": "one"}):
        with pytest.raises(deadlines.RuleError, match="holidays"):
            deadlines.parse_rules({**base, "holidays": [bad]})


# --- county lookup ---------------------------------------------------------------------------------------------

@pytest.mark.parametrize("abbreviation, candidates", [
    ("DCAD", {"Dallas", "Denton"}),
    ("BCAD", {"Bexar", "Brazoria", "Brazos"}),
    ("MCAD", {"Montgomery", "McLennan"}),
])
def test_shared_abbreviations_are_never_resolved_silently(abbreviation, candidates, capsys):
    for extra in ([], ["--case-fragment"], ["--full"]):
        code = county_lookup.main([abbreviation, *extra])
        out = capsys.readouterr().out
        assert code == 1
        assert "matches more than one district" in out and "Nothing was chosen" in out
        assert all(name in out for name in candidates)
        assert '"cad"' not in out and "CAD:" not in out                   # no record, no JSON fragment


def test_an_exact_county_name_beats_somebody_elses_abbreviation(tmp_path, capsys):
    data = {"counties": [{"county": "Alpha", "cad_abbrev": "BETA", "cad_name": "Alpha CAD"},
                         {"county": "Beta", "cad_abbrev": "XCAD", "cad_name": "Beta CAD"}]}
    path = tmp_path / "counties.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    assert county_lookup.main(["beta", "--data", str(path)]) == 0
    assert "Beta County" in capsys.readouterr().out
    assert county_lookup.main(["XCAD", "--data", str(path)]) == 0


@pytest.mark.parametrize("payload", ["[]", '"counties"', '{"counties": {"a": 1}}', '{"counties": ["Harris"]}', "{}"])
def test_county_data_that_is_not_an_object_with_records_is_a_clear_error(payload, tmp_path, capsys):
    path = tmp_path / "counties.json"
    path.write_text(payload, encoding="utf-8")
    assert county_lookup.main(["Harris", "--data", str(path)]) == 2
    assert "expected an object with a \"counties\" list" in capsys.readouterr().err


def test_a_missing_population_rank_does_not_crash_the_list(tmp_path, capsys):
    data = {"counties": [{"county": "Blank", "population_rank": None, "verification_status": None},
                         {"county": "Ranked", "population_rank": 2}]}
    path = tmp_path / "counties.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    assert county_lookup.main(["--list", "--data", str(path)]) == 0
    out = capsys.readouterr().out
    assert " ?  Blank  (?)" in out and " 2  Ranked" in out


def test_a_record_without_a_county_name_does_not_break_not_found(tmp_path, capsys):
    path = tmp_path / "counties.json"
    path.write_text(json.dumps({"counties": [{"cad_abbrev": "ZCAD"}]}), encoding="utf-8")
    assert county_lookup.main(["Nowhere", "--data", str(path)]) == 1
    assert "Covered: ?" in capsys.readouterr().out


def test_docs_mention_the_school_ceiling_and_the_new_rules():
    doc = (SCRIPTS.parent / "references" / "case-schema.md").read_text(encoding="utf-8")
    for fragment in ("`school_ceiling`", "years_after_input", "--acquisition-date", "--qualified-date",
                     "holidays_note", "Gov't Code §662.003(a)", "Other data points (not used in the median)"):
        assert fragment in doc, fragment
