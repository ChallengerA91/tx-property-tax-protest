"""Regression tests for the code-review findings: validation, injection, determinism, atomic writes."""
from __future__ import annotations

import copy
import json
import os
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_package
import deadlines
import savings
from analysis import build_analysis
from case_schema import CaseError, load_case, validate_case
from conftest import EXAMPLES, SCRIPTS, read_json
from deadlines import Deadline
from render_dashboard import BarRow, RefLine, bar_chart

D = deadlines.strict_date


def has_error(errors, fragment):
    return any(fragment in error for error in errors)


# --- savings: capped-homestead low scenario ------------------------------------------------------

def test_capped_low_scenario_applies_the_fraction_to_the_market_reduction_first():
    result = savings.protest_savings(400000, 350000, 2.0, market_value=500000)   # market 500,000, argued 350,000
    assert result.low == 0                       # half of the 150,000 market cut only reaches 425,000: still capped
    assert result.high == pytest.approx(1000.0)  # the whole cut reaches 350,000: 50,000 below the cap


def test_capped_scenarios_cross_the_cap_at_the_right_fraction():
    # market 500,000, appraised 400,000, argued 300,000: a 200,000 cut; 75% of it reaches 350,000
    result = savings.protest_savings(400000, 300000, 2.0, 0.5, 0.75, market_value=500000)
    assert result.low == 0
    assert result.high == pytest.approx(1000.0)


def test_uncapped_scenarios_are_unchanged():
    result = savings.protest_savings(465000, 425000, 2.15)
    assert (result.low, result.high) == (pytest.approx(430.0), pytest.approx(860.0))
    same = savings.protest_savings(465000, 425000, 2.15, market_value=465000)
    assert same == result


# --- validation ------------------------------------------------------------------------------

def test_empty_deadlines_list_does_not_count_as_deadlines(homestead_case):
    homestead_case["deadlines"] = []
    del homestead_case["notice_date"]
    assert has_error(validate_case(homestead_case), "notice_date: provide notice_date")


def test_empty_deadlines_list_with_a_notice_date_computes_them(write_case, homestead_case):
    homestead_case["deadlines"] = []
    a = build_analysis(load_case(write_case(homestead_case)))
    assert [d.id for d in a.deadlines][:2] == ["exemption_application_deadline", "protest_deadline"]


def test_high_fraction_alone_below_the_default_low_is_rejected_up_front(homestead_case):
    homestead_case["settlement_scenarios"] = {"high_fraction": 0.3}
    errors = validate_case(homestead_case)
    assert has_error(errors, "low_fraction (0.5, the default) must not exceed high_fraction (0.3)")


def test_low_fraction_alone_is_fine_up_to_the_default_high(homestead_case):
    homestead_case["settlement_scenarios"] = {"low_fraction": 0.9}
    assert validate_case(homestead_case) == []


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), -float("inf"), 1e30, 10 ** 40])
def test_non_finite_and_absurd_numbers_are_rejected(bad, homestead_case):
    homestead_case["argued_value"] = bad
    homestead_case["comps"][0]["price"] = bad
    homestead_case["subject"]["sqft"] = bad
    errors = validate_case(homestead_case)
    for field in ("argued_value", "comps[0].price", "subject.sqft"):
        assert any(e.startswith(f"{field}: ") and ("finite" in e or "too large" in e) for e in errors), field


def test_json_nan_and_infinity_literals_are_rejected(tmp_path):
    for literal in ("NaN", "Infinity", "-Infinity"):
        path = tmp_path / "case.json"
        path.write_text('{"argued_value": %s}' % literal, encoding="utf-8")
        with pytest.raises(CaseError, match=f"{literal} is not a valid number"):
            load_case(path)


def test_years_are_bounded(homestead_case):
    homestead_case["tax_year"] = 10 ** 6
    homestead_case["comps"][0]["year_built"] = 0
    errors = validate_case(homestead_case)
    assert has_error(errors, "tax_year: must be a year between 1 and 9999")
    assert has_error(errors, "comps[0].year_built: must be a year between 1 and 9999")


def test_net_operating_income_must_be_positive(rental_case):
    rental_case["income"]["expenses"] = [{"label": "Everything", "amount": 40000}]
    errors = validate_case(rental_case)
    assert has_error(errors, "income: net operating income must be above 0")
    assert has_error(errors, "income.gross_rent, income.vacancy_pct and income.expenses")


def test_adjusted_comp_value_must_be_positive(homestead_case):
    homestead_case["comps"][0]["adjustments"] = [{"label": "Huge", "amount": -10 ** 7}]
    errors = validate_case(homestead_case)
    assert has_error(errors, "comps[0]: price plus adjustments must be above 0")
    assert has_error(errors, "comps[0].adjustments")


def test_a_negative_bar_never_gets_a_negative_width():
    svg = bar_chart("c", "t", "d", [BarRow("loss", -5000, "-$5,000"), BarRow("gain", 100, "$100")],
                    [RefLine("ref", -3, "ref-appraised")])
    assert 'width="-' not in svg and 'x="-' not in svg
    assert 'width="0.0"' in svg


@pytest.mark.parametrize("text", ["20270415", "2027-W15-1", "2027-4-5", "2027-04-15T00:00", "2027-04-15 "])
def test_dates_must_be_written_exactly_yyyy_mm_dd(text, homestead_case):
    homestead_case["notice_date"] = text
    assert has_error(validate_case(homestead_case), "notice_date: must be a date written YYYY-MM-DD")


def test_notice_year_must_match_the_tax_year(homestead_case):
    homestead_case["notice_date"] = "2026-04-01"
    assert has_error(validate_case(homestead_case), "notice_date: year 2026 must equal tax_year 2027")


def test_prepared_on_may_not_precede_the_notice(homestead_case):
    homestead_case["prepared_on"] = "2027-03-01"
    assert has_error(validate_case(homestead_case), "prepared_on: 2027-03-01 is earlier than notice_date 2027-04-01")


def test_school_rate_may_not_exceed_the_total_rate(homestead_case):
    homestead_case["tax_rates"]["school"] = 3.0
    assert has_error(validate_case(homestead_case), "tax_rates.school: must not exceed tax_rates.total (2.15)")


def test_text_fields_reject_control_characters(homestead_case):
    homestead_case["subject"]["address"] = "12 Main\x07 St"
    homestead_case["comps"][0]["notes"] = "line\rbreak"
    homestead_case["market_notes"] = ["fine\ttab\nand newline"]
    errors = validate_case(homestead_case)
    assert has_error(errors, "subject.address: contains a control character (code 0x07)")
    assert has_error(errors, "comps[0].notes: contains a control character (code 0x0d)")
    assert not has_error(errors, "market_notes")


def test_case_files_may_start_with_a_byte_order_mark(tmp_path, homestead_case):
    path = tmp_path / "case.json"
    path.write_bytes(b"\xef\xbb\xbf" + json.dumps(homestead_case).encode("utf-8"))
    assert load_case(path)["path"] == "homestead"


def test_non_utf8_case_files_get_a_clear_error(tmp_path):
    path = tmp_path / "case.json"
    path.write_bytes(b'{"county": "Caf\xe9"}')
    with pytest.raises(CaseError, match="not UTF-8 text"):
        load_case(path)


# --- savings inputs --------------------------------------------------------------------------

@pytest.mark.parametrize("case, message", [
    ({"subject": {"appraised_value": 100}, "argued_value": -5, "tax_rates": {"total": 2}}, "argued_value"),
    ({"subject": {"appraised_value": 100}, "argued_value": 50, "tax_rates": {"total": -2}}, "tax_rates.total"),
    ({"subject": {"appraised_value": 100}, "argued_value": 50, "tax_rates": {"total": float("nan")}}, "finite"),
    ({"subject": {"appraised_value": 100}, "argued_value": 50, "tax_rates": {"total": 500}}, "tax_rates.total"),
    ({"subject": {"appraised_value": 100}, "argued_value": 50, "tax_rates": {"total": 2, "school": 3}}, "must not exceed"),
    ({"subject": {"appraised_value": 100}, "argued_value": 100, "tax_rates": {"total": 2}}, "must be below the market value"),
    ({"subject": {"appraised_value": 100, "market_value": 90}, "argued_value": 50, "tax_rates": {"total": 2}}, "subject.market_value"),
    ({"subject": {"appraised_value": float("inf")}, "argued_value": 50, "tax_rates": {"total": 2}}, "finite"),
])
def test_savings_rejects_inputs_that_could_exceed_the_tax_bill(case, message):
    with pytest.raises(ValueError, match=message):
        savings.estimate_savings(case)


def test_savings_cli_validates_flags(capsys):
    assert savings.main(["--appraised", "100000", "--argued", "-1", "--total-rate", "2"]) == 2
    assert "argued_value" in capsys.readouterr().err
    assert savings.main(["--appraised", "400000", "--market", "500000", "--argued", "450000",
                         "--total-rate", "2", "--meter"]) == 0
    assert "No protest savings this year" in capsys.readouterr().out


@pytest.mark.parametrize("appraised, market, argued", [(400000, 500000, 1), (400000, 500000, 399999),
                                                       (400000, 400000, 200000), (400000, 900000, 450000)])
def test_total_savings_never_exceed_the_tax_bill(appraised, market, argued):
    case = {"subject": {"appraised_value": appraised, "market_value": market}, "argued_value": argued,
            "tax_rates": {"total": 2.0, "school": 1.5},
            "exemptions": {"eligible": [{"name": "School", "amount": 10 ** 9, "applies_to": "school"},
                                        {"name": "Full", "kind": "total"}]}}
    estimate = savings.estimate_savings(case)
    assert estimate.total.high <= savings.tax_on(appraised, 2.0) + 0.01


# --- iCalendar ------------------------------------------------------------------------------

STAMP = deadlines.dt.datetime(2027, 4, 20, 12, 0, 0, tzinfo=deadlines.dt.timezone.utc)


def unfold(text):
    return text.replace("\r\n ", "").split("\r\n")


def test_ics_text_cannot_break_out_with_a_bare_cr_or_control_characters():
    item = Deadline("x", "Event\rEND:VEVENT\rBEGIN:VEVENT\x07", D("2027-06-01"), D("2027-06-01"),
                    "Tax\x00Code", "note\rwith CR\r\nand CRLF\x1b")
    text = deadlines.build_ics([item], 2027, STAMP)
    assert "\r" not in text.replace("\r\n", "")
    assert not any(ch in text for ch in "\x00\x07\x1b")
    lines = unfold(text)
    assert lines.count("BEGIN:VEVENT") == 1 and lines.count("END:VEVENT") == 1
    assert "SUMMARY:Event\\nEND:VEVENT\\nBEGIN:VEVENT" in lines


def test_events_with_names_that_normalize_the_same_get_distinct_uids(write_case, homestead_case):
    homestead_case["deadlines"] = [{"event": "Hearing", "date": "2027-06-01"},
                                   {"event": "hearing!", "date": "2027-06-02"},
                                   {"event": "HEARING", "date": "2027-06-03"}]
    a = build_analysis(load_case(write_case(homestead_case)))
    assert [d.id for d in a.deadlines] == ["hearing", "hearing-2", "hearing-3"]
    uids = [line for line in unfold(deadlines.build_ics(a.deadlines, 2027, STAMP)) if line.startswith("UID:")]
    assert len(set(uids)) == 3
    again = [line for line in unfold(deadlines.build_ics(a.deadlines, 2027, STAMP)) if line.startswith("UID:")]
    assert uids == again


def test_build_ics_makes_repeated_ids_unique_and_stable():
    items = [Deadline("same", "A", D("2027-06-01"), D("2027-06-01"), "", ""),
             Deadline("same", "B", D("2027-06-02"), D("2027-06-02"), "", "")]
    uids = [line for line in unfold(deadlines.build_ics(items, 2027, STAMP)) if line.startswith("UID:")]
    assert len(set(uids)) == 2
    assert uids == [line for line in unfold(deadlines.build_ics(items, 2027, STAMP)) if line.startswith("UID:")]


def test_a_naive_dtstamp_is_taken_as_utc():
    naive = deadlines.dt.datetime(2027, 4, 20, 12, 30, 5)
    item = Deadline("x", "E", D("2027-06-01"), D("2027-06-01"), "", "")
    assert "DTSTAMP:20270420T123005Z" in unfold(deadlines.build_ics([item], 2027, naive))


# --- injection: spreadsheet, markdown, HTML ------------------------------------------------------

FORMULAS = ['=HYPERLINK("http://evil.example/","click")', "=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(1,1)",
            "\t=1+1"]


@pytest.mark.parametrize("payload", FORMULAS)
def test_spreadsheet_text_never_becomes_a_formula(payload, tmp_path, write_case, homestead_case, rental_case):
    for case in (homestead_case, rental_case):
        case["subject"]["address"] = payload
        case["comps"][0]["address"] = payload
        case["comps"][0]["notes"] = payload
        case["market_notes"] = [payload]
        out = tmp_path / case["path"]
        build_package.build_package(write_case(case, f"{case['path']}.json"), out)
        path = out / "comp_analysis.xlsx"
        workbook = load_workbook(path)
        found = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and payload in cell.value:
                        found += 1
                        assert cell.data_type == "s", (sheet.title, cell.coordinate)
        assert found >= 3
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/"):
                    xml = archive.read(name).decode("utf-8")
                    assert "<f>" not in xml and "<f " not in xml, name


def test_ordinary_text_still_round_trips_through_the_workbook(tmp_path, write_case, homestead_case):
    build_package.build_package(write_case(homestead_case), tmp_path / "out")
    sheet = load_workbook(tmp_path / "out" / "comp_analysis.xlsx")["Comp Analysis"]
    assert any(c.value == homestead_case["subject"]["address"] for row in sheet.iter_rows() for c in row)


XSS = "<img src=x onerror=alert(1)>\"'&<script>alert(2)</script>\nEND:VEVENT"
NOT_FREE_TEXT = {"path", "strategy", "price_type", "applies_to", "kind"}
DATE_TEXT = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def inject(value, key=""):
    """The same case with an attack string appended to every free-text field."""
    if isinstance(value, dict):
        return {k: (v if k.startswith("_") else inject(v, k)) for k, v in value.items()}
    if isinstance(value, list):
        return [inject(v, key) for v in value]
    if isinstance(value, str) and key not in NOT_FREE_TEXT and not DATE_TEXT.match(value):
        return value + XSS
    return value


class TagCollector(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.attributes = [], []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attributes += [name for name, _ in attrs]


ALLOWED_TAGS = {"html", "head", "meta", "title", "style", "body", "a", "header", "p", "h1", "h2", "h3", "button",
                "main", "section", "dl", "div", "dt", "dd", "span", "br", "figure", "figcaption", "svg", "desc",
                "line", "text", "rect", "table", "caption", "thead", "tbody", "tfoot", "tr", "th", "td", "details",
                "summary", "ul", "li", "ol", "time", "strong", "footer", "script"}


@pytest.mark.parametrize("sample", ["sample-case.json", "sample-case-rental.json", "sample-case-capped.json"])
def test_attack_text_in_every_free_text_field_stays_inert(sample, tmp_path, write_case):
    case = inject(read_json(EXAMPLES / sample))
    assert validate_case(case) == []
    out = tmp_path / "out"
    build_package.build_package(write_case(case), out)

    parser = TagCollector()
    parser.feed((out / "dashboard.html").read_text(encoding="utf-8"))
    assert set(parser.tags) <= ALLOWED_TAGS
    assert parser.tags.count("script") == 1 and "img" not in parser.tags
    assert not [name for name in parser.attributes if name.startswith("on")]

    checklist = (out / "filing_checklist.md").read_text(encoding="utf-8")
    assert "<" not in checklist and ">" not in checklist
    assert "&lt;img src=x onerror=alert(1)&gt;" in checklist
    for line in checklist.splitlines():
        if line.startswith("|"):
            assert line.endswith("|")          # a table row never spills onto a second line

    ics = unfold((out / "deadlines.ics").read_bytes().decode("utf-8"))
    assert ics.count("BEGIN:VEVENT") == ics.count("END:VEVENT") == len(build_analysis(case).deadlines)
    assert all(re.match(r"^[A-Z][A-Z-]*(;[^:]*)?:", line) for line in ics if line)


def test_markdown_cells_cannot_break_the_table_or_inject_markup(tmp_path, write_case, homestead_case):
    homestead_case["exemptions"]["eligible"][0]["name"] = "Over-65\nspecial | with pipe <b>bold</b> & [link](http://evil.example)"
    build_package.build_package(write_case(homestead_case), tmp_path / "out")
    text = (tmp_path / "out" / "filing_checklist.md").read_text(encoding="utf-8")
    rows = [line for line in text.splitlines() if "Exemption: Over-65" in line]
    assert len(rows) == 1
    assert "Over-65 special \\| with pipe &lt;b&gt;bold&lt;/b&gt; &amp; \\[link\\](http://evil.example)" in rows[0]
    assert "<b>" not in text


def test_dashboard_carries_a_restrictive_content_security_policy(tmp_path, write_case, homestead_case):
    build_package.build_package(write_case(homestead_case), tmp_path / "out")
    html = (tmp_path / "out" / "dashboard.html").read_text(encoding="utf-8")
    assert ('<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; '
            'script-src \'unsafe-inline\'; img-src data:">') in html
    assert html.index("Content-Security-Policy") < html.index("<style>")


def test_year_built_is_escaped_in_the_dashboard():
    from render_dashboard import _comp_row
    from analysis import build_comp_rows
    row = build_comp_rows([{"address": "1 A", "price": 1, "year_built": 1999}])[0]
    assert "<td class=\"num\">1999</td>" in _comp_row(row)


# --- determinism and packaging ---------------------------------------------------------------

def test_ooxml_files_are_stored_with_one_creating_system_and_a_fixed_date(tmp_path, write_case, homestead_case):
    build_package.build_package(write_case(homestead_case), tmp_path / "out")
    for name in ("comp_analysis.xlsx", "protest_letter.docx"):
        with zipfile.ZipFile(tmp_path / "out" / name) as archive:
            infos = archive.infolist()
            assert [i.filename for i in infos] == sorted(i.filename for i in infos)
            for info in infos:
                assert info.create_system == 3 and info.compress_type == zipfile.ZIP_STORED
                assert info.date_time == (1980, 1, 1, 0, 0, 0)
            assert archive.testzip() is None


def test_requirements_pin_the_libraries_and_keep_pytest_out():
    runtime = (SCRIPTS / "requirements.txt").read_text(encoding="utf-8")
    dev = (SCRIPTS / "requirements-dev.txt").read_text(encoding="utf-8")
    assert "pytest" not in runtime
    assert re.search(r"openpyxl>=3\.1\.2,<3\.2", runtime) and re.search(r"python-docx>=1\.1\.0,<2", runtime)
    assert "-r requirements.txt" in dev and re.search(r"pytest>=7", dev)


def test_the_determinism_claim_names_its_conditions():
    for text in (build_package.__doc__, (SCRIPTS.parent / "references" / "case-schema.md").read_text(encoding="utf-8")):
        flat = " ".join(text.split())
        assert "same Python and library versions" in flat and "byte-identical" in flat


def test_missing_prepared_on_uses_a_neutral_placeholder_everywhere(tmp_path, write_case, homestead_case):
    del homestead_case["prepared_on"]
    out = tmp_path / "out"
    build_package.build_package(write_case(homestead_case), out)
    assert "DTSTAMP:19800101T000000Z" in (out / "deadlines.ics").read_bytes().decode("utf-8")
    with zipfile.ZipFile(out / "comp_analysis.xlsx") as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert "1980-01-01T00:00:00Z" in core and "2000" not in core
    import docx
    created = docx.Document(out / "protest_letter.docx").core_properties.created
    assert (created.year, created.month, created.day) == (1980, 1, 1)


# --- atomic publishing ---------------------------------------------------------------------------

def package_files(folder: Path):
    return {p.name: p.read_bytes() for p in folder.iterdir()}


def test_a_failed_move_leaves_the_existing_package_untouched(tmp_path, write_case, homestead_case, monkeypatch):
    out = tmp_path / "out"
    build_package.build_package(write_case(homestead_case, "first.json"), out)
    before = package_files(out)
    changed = copy.deepcopy(homestead_case)
    changed["argued_value"] = 430000
    real_replace = os.replace

    def locked(src, dst, *args, **kwargs):
        if Path(dst).name == "dashboard.html" and ".staging-" in str(src):
            raise PermissionError(13, "The process cannot access the file because it is being used")
        return real_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(PermissionError):
        build_package.build_package(write_case(changed, "second.json"), out)
    monkeypatch.undo()
    assert package_files(out) == before                       # every original file, byte for byte, and no leftovers
    assert sorted(before) == sorted(build_package.OUTPUT_NAMES.values())


def test_a_failed_first_build_leaves_no_partial_package(tmp_path, write_case, homestead_case, monkeypatch):
    out = tmp_path / "out"
    real_replace = os.replace

    def locked(src, dst, *args, **kwargs):
        if Path(dst).name == "deadlines.ics" and ".staging-" in str(src):
            raise PermissionError(13, "locked")
        return real_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(PermissionError):
        build_package.build_package(write_case(homestead_case), out)
    assert list(out.iterdir()) == []


def test_cli_explains_a_write_failure_in_plain_words(tmp_path, write_case, homestead_case, monkeypatch, capsys):
    real_replace = os.replace

    def locked(src, dst, *args, **kwargs):
        if Path(dst).name == "comp_analysis.xlsx" and ".staging-" in str(src):
            raise PermissionError(13, "Permission denied")
        return real_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", locked)
    code = build_package.main([str(write_case(homestead_case)), "--out", str(tmp_path / "out")])
    assert code == 1
    err = capsys.readouterr().err
    assert "Could not write the package" in err and "Permission denied" in err and "nothing was changed" in err
